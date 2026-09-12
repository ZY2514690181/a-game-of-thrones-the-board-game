# 上次更新：The Board于北京时间2026年9月3日

import os
import random

import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font


# Parameters
NUM_OF_PLAYERS = 24
RANDOM_SEED_INPUTS = [0, 42, 1]

script_dir = os.path.dirname(os.path.abspath(__file__))

# Input files
PLAYER_NAMES_FILE = os.path.join(script_dir, "Input files", "选手名.xlsx")

# Output files
PLAYER_IDS_FILE = os.path.join(script_dir, "Output files", "选手号码分配.xlsx")
TOURNAMENT_SCHEDULE_PLAYER_IDS_FILE = os.path.join(script_dir, "Output files", "比赛排表结果（选手号码）.xlsx")
TOURNAMENT_SCHEDULE_PLAYER_NAMES_FILE = os.path.join(script_dir, "Output files", "比赛排表结果（选手名）.xlsx")
PLAYER_MESSAGES_FILE = os.path.join(
    script_dir,
    "Output files",
    "私发选手消息.txt"
)

# 6 colors of houses in tournament
HOUSE_COLORS = [
    "白",
    "黑",
    "红",
    "黄",
    "绿",
    "橙"
]
reordered_house_colors = HOUSE_COLORS.copy()


# Computation results are provided by 小X
FIRST_ROW_TEMPLATE = {
    6: [1,2,3,4,5,6],
    7: [1,2,3,4,5,6],
    8: [1,2,3,4,5,6],
    9: [1,2,3,4,5,7],
    10: [1,2,3,4,6,7],
    11: [1,2,3,5,6,8],
    12: [1,2,3,4,6,9],
    13: [1,2,3,4,6,10],
    14: [1,2,3,4,6,10],
    15: [1,2,3,4,7,11],
    18: [1,2,3,5,9,14],
    19: [1,2,3,5,8,12],
    20: [1,2,3,5,8,13],
    21: [1,2,3,5,8,13],
    22: [1,2,3,5,9,14],
    23: [1,2,3,5,8,16],
    24: [1,2,3,5,13,20],
    25: [1,2,3,5,10,16],
    26: [1,2,3,6,10,16],
    27: [1,2,3,6,14,23],
    28: [1,2,5,16,21,23],
    31: [1,2,4,9,13,19],
    35: [1,2,4,8,13,21],
    36: [1,2,4,9,24,28],
    37: [1,2,4,8,17,27],
    38: [1,2,4,8,18,31],
    39: [1,2,4,8,13,23],
    40: [1,2,4,8,13,21]
}


def get_template_player_count():
    """
    Return the template size used for generating the tournament board.

    If the number of players exceeds 40, use the 40-player template.
    """
    if NUM_OF_PLAYERS > 40:
        return 40

    return NUM_OF_PLAYERS


def validate_configuration():
    """Validate user-configurable parameters."""

    if not isinstance(NUM_OF_PLAYERS, int):
        raise ValueError("NUM_OF_PLAYERS 必须为整数。")

    if NUM_OF_PLAYERS < 6:
        raise ValueError("NUM_OF_PLAYERS 不得小于 6。")

    unsupported_player_counts = {16, 17, 29, 30, 32, 33, 34}
    if NUM_OF_PLAYERS in unsupported_player_counts:
        raise ValueError(
            f"目前暂不支持 {NUM_OF_PLAYERS} 名选手。"
        )

    template_player_count = get_template_player_count()

    if template_player_count not in FIRST_ROW_TEMPLATE:
        raise ValueError(
            f"未找到 {template_player_count} 名选手对应的比赛模板。"
        )

    if not isinstance(RANDOM_SEED_INPUTS, list):
        raise ValueError(
            "RANDOM_SEED_INPUTS 必须为列表。"
        )

    if len(RANDOM_SEED_INPUTS) == 0:
        raise ValueError(
            "RANDOM_SEED_INPUTS 不得为空。"
        )

    for i, seed in enumerate(RANDOM_SEED_INPUTS, start=1):
        if not isinstance(seed, int):
            raise ValueError(
                f"第 {i} 个随机种子必须为整数。"
            )


def load_player_names():
    """
    Load player names Excel file.
    """

    try:
        player_names_df = pd.read_excel(
            PLAYER_NAMES_FILE,
            header=None
        )

    except FileNotFoundError:
        raise FileNotFoundError(
            f"找不到选手名单文件：\n{PLAYER_NAMES_FILE}"
        )

    except PermissionError:
        raise PermissionError(
            "无法读取选手名单文件，请关闭 Excel 后重试。"
        )

    except Exception as e:
        raise Exception(
            f"选手名单文件读取失败：{e}"
        )

    return player_names_df


def validate_player_names(player_names_df):
    """
    Validate player aliases input file.
    """

    # Check only one column
    if player_names_df.shape[1] != 1:
        raise ValueError(
            "选手名.xlsx 必须只包含一列选手名。"
        )

    # Extract player aliases
    player_names = player_names_df.iloc[:, 0]

    # Check blank Excel cells
    if player_names.isna().any():
        raise ValueError(
            "选手名单中存在空白行。"
        )

    # Convert all values to string
    player_names = player_names.astype(str)

    # Check empty strings only
    # Note: spaces are considered valid aliases
    if (player_names == "").any():
        raise ValueError(
            "选手名单中存在空选手名。"
        )

    # Check exact duplicate aliases
    duplicated_names = player_names[player_names.duplicated()].unique()

    if len(duplicated_names) > 0:
        raise ValueError(
            f"选手名单中存在重复选手名：{list(duplicated_names)}"
        )

    # Check player count
    if len(player_names) != NUM_OF_PLAYERS:
        raise ValueError(
            f"NUM_OF_PLAYERS={NUM_OF_PLAYERS}，"
            f"但选手名单中共有 {len(player_names)} 人。"
        )

    return player_names.tolist()


def initialize_random_seed():
    """
    Initialize random generator using combined judge seeds.
    """

    random_seed = sum(RANDOM_SEED_INPUTS)
    random.seed(random_seed)


def generate_player_numbers(player_names):
    """
    Randomly assign player numbers.

    Return:
        player_name_to_id: dict
        player_id_to_name: dict
    """

    player_numbers = list(range(1, NUM_OF_PLAYERS + 1))

    random.shuffle(player_numbers)

    player_name_to_id = {}
    player_id_to_name = {}

    for name, number in zip(player_names, player_numbers):
        player_name_to_id[name] = number
        player_id_to_name[number] = name

    return player_name_to_id, player_id_to_name


def generate_schedule_player_ids():
    """
    Generate tournament schedule using player numbers.

    Return:
        pandas DataFrame
    """

    template_size = get_template_player_count()

    # Generate shuffled game numbers
    game_numbers = list(range(1, NUM_OF_PLAYERS + 1))
    random.shuffle(game_numbers)

    # Column names
    columns = ["局号"] + reordered_house_colors

    rows = []

    # First row
    first_players = FIRST_ROW_TEMPLATE[template_size]

    rows.append(
        [
            game_numbers[0]
        ]
        + first_players
    )

    # Remaining rows
    for i in range(1, NUM_OF_PLAYERS):

        previous_players = rows[-1][1:]

        next_players = [
            (player_number % NUM_OF_PLAYERS) + 1
            for player_number in previous_players
        ]

        rows.append(
            [
                game_numbers[i]
            ]
            + next_players
        )

    return pd.DataFrame(
        rows,
        columns=columns
    )


def validate_schedule(schedule_player_ids):
    """
    Validate generated tournament schedule.
    """

    validate_schedule_columns(schedule_player_ids)
    validate_schedule_player_counts(schedule_player_ids)
    validate_schedule_rows(schedule_player_ids)


def validate_schedule_columns(schedule_player_ids):
    """
    Verify every color column is a permutation of 1..NUM_OF_PLAYERS.
    """

    expected_numbers = list(range(1, NUM_OF_PLAYERS + 1))

    for column in HOUSE_COLORS:
        actual_numbers = sorted(schedule_player_ids[column].tolist())

        if actual_numbers != expected_numbers:
            raise ValueError(
                f"比赛排表验证失败：'{column}'列不是 1 至 {NUM_OF_PLAYERS} 的一个排列。"
            )


def validate_schedule_player_counts(schedule_player_ids):
    """
    Verify every player appears exactly six times.
    """

    player_counts = {}

    for column in HOUSE_COLORS:
        for player in schedule_player_ids[column]:
            player_counts[player] = player_counts.get(player, 0) + 1

    for player in range(1, NUM_OF_PLAYERS + 1):
        if player_counts.get(player, 0) != 6:
            raise ValueError(
                f"比赛排表验证失败：选手 {player} 共出现 "
                f"{player_counts.get(player, 0)} 次，应为 6 次。"
            )


def validate_schedule_rows(schedule_player_ids):
    """
    Verify every game contains six distinct players.
    """

    for _, row in schedule_player_ids.iterrows():

        players = row[HOUSE_COLORS].tolist()

        if len(set(players)) != 6:
            raise ValueError(
                f"比赛排表验证失败：第 {row['局号']} 局存在重复选手。"
            )


def generate_schedule_player_names(
    schedule_player_ids,
    player_id_to_name
):
    """
    Replace player numbers with names.

    Return:
        pandas DataFrame sorted by 局号.
    """

    schedule_player_names = schedule_player_ids.copy()

    for column in HOUSE_COLORS:
        schedule_player_names[column] = (
            schedule_player_names[column]
            .map(player_id_to_name)
        )

    schedule_player_names = (
        schedule_player_names
        .loc[:, ["局号"] + HOUSE_COLORS]
        .sort_values("局号")
        .reset_index(drop=True)
    )

    return schedule_player_names


def build_player_schedule(schedule_player_ids):
    """
    Build player schedule.

    Return:
        {
            player_number: {
                "白": game_number,
                ...
            }
        }
    """

    player_schedule = {}

    for _, row in schedule_player_ids.iterrows():

        game_number = row["局号"]

        for color in HOUSE_COLORS:

            player_number = row[color]

            if player_number not in player_schedule:
                player_schedule[player_number] = {}

            player_schedule[player_number][color] = game_number

    return player_schedule


def ensure_output_directory():
    """
    Create output directory if it does not exist.
    """

    output_directory = os.path.dirname(PLAYER_IDS_FILE)

    if not os.path.exists(output_directory):
        os.makedirs(output_directory)


def format_excel_workbook(excel_file):
    """
    Apply formatting to an Excel workbook.
    """

    workbook = load_workbook(excel_file)

    font = Font(name="宋体")

    alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    for worksheet in workbook.worksheets:

        # Apply font and alignment
        for row in worksheet.iter_rows():
            for cell in row:
                cell.font = font
                cell.alignment = alignment

        # Auto-adjust column widths
        for column_cells in worksheet.columns:

            max_length = 0

            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[
                column_cells[0].column_letter
            ].width = max_length + 10

    workbook.save(excel_file)


def save_player_numbers(player_name_to_id):
    """
    Save player name-number assignment.
    """

    ensure_output_directory()

    player_numbers_df = pd.DataFrame(
        list(player_name_to_id.items()),
        columns=["选手名", "号码"]
    )

    try:
        player_numbers_df.to_excel(
            PLAYER_IDS_FILE,
            index=False
        )
        format_excel_workbook(
            PLAYER_IDS_FILE
        )

    except PermissionError:
        raise PermissionError(
            f"无法写入文件，请关闭 Excel 后重试：\n"
            f"{PLAYER_IDS_FILE}"
        )

    except Exception as e:
        raise Exception(
            f"保存选手号码分配文件失败：{e}"
        ) from e


def save_schedule_player_ids(schedule_player_ids):
    """
    Save tournament schedule using player numbers.
    """

    ensure_output_directory()

    try:
        schedule_player_ids.to_excel(
            TOURNAMENT_SCHEDULE_PLAYER_IDS_FILE,
            index=False
        )
        format_excel_workbook(
            TOURNAMENT_SCHEDULE_PLAYER_IDS_FILE
        )

    except PermissionError:
        raise PermissionError(
            f"无法写入文件，请关闭 Excel 后重试：\n"
            f"{TOURNAMENT_SCHEDULE_PLAYER_IDS_FILE}"
        )

    except Exception as e:
        raise Exception(
            f"保存号码比赛排表失败：{e}"
        )


def save_schedule_player_names(schedule_player_names):
    """
    Save tournament schedule using player names.
    """

    ensure_output_directory()

    try:
        schedule_player_names.to_excel(
            TOURNAMENT_SCHEDULE_PLAYER_NAMES_FILE,
            index=False
        )
        format_excel_workbook(
            TOURNAMENT_SCHEDULE_PLAYER_NAMES_FILE
        )

    except PermissionError:
        raise PermissionError(
            f"无法写入文件，请关闭 Excel 后重试：\n"
            f"{TOURNAMENT_SCHEDULE_PLAYER_NAMES_FILE}"
        )

    except Exception as e:
        raise Exception(
            f"保存选手名比赛排表失败：{e}"
        )


def save_player_messages(
    player_names,
    player_name_to_id,
    player_schedule
):
    """
    Save private player messages.
    """

    ensure_output_directory()

    try:

        with open(
            PLAYER_MESSAGES_FILE,
            "w",
            encoding="utf-8-sig"
        ) as file:

            for i, player_name in enumerate(player_names):

                player_number = player_name_to_id[player_name]

                file.write(
                    f"【请私发给选手{player_name}】\n"
                )

                file.write(
                    f"选手{player_name}你好，你的比赛局号为：\n"
                )

                for color in HOUSE_COLORS:

                    game_number = player_schedule[player_number][color]

                    file.write(
                        f"{color}——局{game_number}\n"
                    )

                if i != len(player_names) - 1:
                    file.write("\n")

    except PermissionError:
        raise PermissionError(
            f"无法写入文件，请关闭文本文件后重试：\n"
            f"{PLAYER_MESSAGES_FILE}"
        )

    except Exception as e:
        raise Exception(
            f"保存私发选手消息失败：{e}"
        ) from e


def main():
    # Phase 1: Validate configuration
    validate_configuration()

    # Phase 2: Load files
    player_names_df = load_player_names()

    # Phase 3: Validate input
    player_names = validate_player_names(player_names_df)

    # Phase 4: Generate data
    initialize_random_seed()

    random.shuffle(reordered_house_colors)

    player_name_to_id, player_id_to_name = (
        generate_player_numbers(player_names)
    )

    schedule_player_ids = (
        generate_schedule_player_ids()
    )

    validate_schedule(schedule_player_ids)

    schedule_player_names = (
        generate_schedule_player_names(
            schedule_player_ids,
            player_id_to_name
        )
    )

    player_schedule = (
        build_player_schedule(
            schedule_player_ids
        )
    )

    # Phase 5: Save outputs
    save_player_numbers(
        player_name_to_id
    )

    save_schedule_player_ids(
        schedule_player_ids
    )

    save_schedule_player_names(
        schedule_player_names
    )

    save_player_messages(
        player_names,
        player_name_to_id,
        player_schedule
    )

    print("比赛排表生成完成。")
    print(f"输出文件位置：{os.path.dirname(PLAYER_IDS_FILE)}")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"\n错误：{e}")
        input("\n按 Enter 键退出...")
